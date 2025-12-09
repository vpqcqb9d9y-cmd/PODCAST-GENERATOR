from __future__ import annotations

import json
import shutil
import textwrap
from pathlib import Path
from typing import Dict, Iterable, List, Optional

import requests
from openai import AzureOpenAI
from PyPDF2 import PdfReader
from pptx import Presentation
from docx import Document

from ..utils import Settings, get_logger

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".md"}


class MetadataIngestor:
    """Summarize supporting materials (files/URLs) and merge into metadata."""

    def __init__(self, settings: Settings) -> None:
        self.logger = get_logger(self.__class__.__name__)
        self.settings = settings
        self.client = AzureOpenAI(
            api_key=settings.openai_api_key,
            api_version=settings.openai_api_version,
            azure_endpoint=settings.openai_endpoint,
        )

    def enrich(
        self,
        metadata: Dict,
        material_paths: Iterable[Path],
        urls: Iterable[str],
        run_paths,
    ) -> Dict:
        summaries: List[Dict] = []
        material_paths = list(material_paths)
        urls = [u.strip() for u in urls if u.strip()]

        if not material_paths and not urls:
            return metadata

        if material_paths:
            run_paths.materials_dir.mkdir(parents=True, exist_ok=True)
        for path in material_paths:
            if not path.exists():
                self.logger.warning("Material file %s not found, skipping.", path)
                continue
            text = self._extract_text(path)
            if not text:
                self.logger.warning("No extractable text from %s", path.name)
                continue
            summary = self._summarize(text, title=path.name)
            summary["source_path"] = str(path)
            summaries.append(summary)
            try:
                shutil.copy(path, run_paths.materials_dir / path.name)
            except OSError as exc:
                self.logger.warning("Failed to copy %s into run folder: %s", path, exc)

        for url in urls:
            text = self._download_text(url)
            if text:
                summary = self._summarize(text, title=url)
                summary["source_url"] = url
                summaries.append(summary)
            else:
                self.logger.warning("Failed to fetch %s", url)

        if not summaries:
            return metadata

        metadata.setdefault("supporting_materials", [])
        metadata["supporting_materials"].extend(summaries)

        labs = metadata.setdefault("labs", [])
        references = metadata.setdefault("references", [])
        reading = metadata.setdefault("reading_list", [])

        for summary in summaries:
            for lab in summary.get("labs", []):
                if lab not in labs:
                    labs.append(lab)
            for ref in summary.get("references", []):
                if ref not in references:
                    references.append(ref)
            for item in summary.get("reading_list", []):
                if item not in reading:
                    reading.append(item)

        return metadata

    def _extract_text(self, path: Path) -> str:
        ext = path.suffix.lower()
        if ext == ".pdf":
            return self._text_from_pdf(path)
        if ext == ".docx":
            return self._text_from_docx(path)
        if ext == ".pptx":
            return self._text_from_pptx(path)
        if ext in {".txt", ".md"}:
            return path.read_text(encoding="utf-8", errors="ignore")
        self.logger.warning("Unsupported extension %s (defaulting to txt)", ext)
        return path.read_text(encoding="utf-8", errors="ignore")

    def _text_from_pdf(self, path: Path) -> str:
        try:
            reader = PdfReader(str(path))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as exc:
            self.logger.warning("Failed to read PDF %s: %s", path.name, exc)
            return ""

    def _text_from_docx(self, path: Path) -> str:
        try:
            doc = Document(str(path))
            paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
            return "\n".join(paragraphs)
        except Exception as exc:
            self.logger.warning("Failed to read DOCX %s: %s", path.name, exc)
            return ""

    def _text_from_pptx(self, path: Path) -> str:
        try:
            presentation = Presentation(str(path))
            texts = []
            for idx, slide in enumerate(presentation.slides, start=1):
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text.strip())
                if slide_text:
                    texts.append(f"Slide {idx}: " + " | ".join(slide_text))
            return "\n".join(texts)
        except Exception as exc:
            self.logger.warning("Failed to read PPTX %s: %s", path.name, exc)
            return ""

    def _download_text(self, url: str) -> str:
        try:
            resp = requests.get(url, timeout=10)
            resp.raise_for_status()
            return resp.text
        except requests.RequestException as exc:
            self.logger.warning("Failed to download %s: %s", url, exc)
            return ""

    def _summarize(self, text: str, title: str) -> Dict:
        truncated = textwrap.shorten(text, width=8000, placeholder=" ...")
        prompt = f"""You are an assistant that extracts actionable study metadata from supporting materials for Azure lectures.
Return JSON with keys:
- "title": short title
- "summary": concise Hebrew summary
- "labs": list of lab exercises (strings)
- "references": list of links or papers
- "reading_list": list of recommended follow-up items

Material title: {title}
Material content:
{truncated}
"""
        response = self.client.chat.completions.create(
            model=self.settings.openai_deployment,
            messages=[
                {"role": "system", "content": "You summarize Azure lecture materials into structured JSON."},
                {"role": "user", "content": prompt},
            ],
            temperature=0.3,
            max_tokens=600,
            response_format={"type": "json_object"},
        )
        content = response.choices[0].message.content
        try:
            payload = json.loads(content)
        except json.JSONDecodeError:
            payload = {"title": title, "summary": content}
        payload.setdefault("title", title)
        payload.setdefault("summary", "")
        payload.setdefault("labs", [])
        payload.setdefault("references", [])
        payload.setdefault("reading_list", [])
        return payload

